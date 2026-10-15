"""
Rutas para exportar ganancias a Excel

Este módulo maneja la generación y descarga de reportes
en formato Excel (.xlsx) con los datos de ventas.
"""

from fastapi import APIRouter, Response
from fastapi.responses import StreamingResponse
from models import ExportRequest
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import datetime

router = APIRouter()


@router.post("/export")
async def export_to_excel(request: ExportRequest):
    """
    Genera un archivo Excel con el reporte de ganancias
    
    Recibe una lista de ventas desde el frontend y las convierte
    en un archivo Excel formateado en español.
    """
    
    # Crear nuevo libro de Excel
    wb = Workbook()
    
    # ==========================================
    # HOJA 1: Resumen General
    # ==========================================
    ws_resumen = wb.active
    ws_resumen.title = "Resumen"
    
    # Estilos
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_fill = PatternFill(start_color="16A34A", end_color="16A34A", fill_type="solid")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Calcular totales
    total_ganancias = sum(v.precioVenta for v in request.ventas)
    ventas_lucio = [v for v in request.ventas if v.tipoVenta == "LUCIO"]
    ventas_doctores = [v for v in request.ventas if v.tipoVenta == "DOCTORES"]
    ventas_envio = [v for v in request.ventas if v.tipoVenta == "ENVIO"]
    
    total_qr = sum(v.precioVenta for v in ventas_lucio if v.metodoPago == "QR")
    total_cash = sum(v.precioVenta for v in ventas_lucio if v.metodoPago == "EFECTIVO")
    
    # Título
    ws_resumen["A1"] = "REPORTE DE GANANCIAS"
    ws_resumen["A1"].font = Font(bold=True, size=16)
    ws_resumen.merge_cells("A1:D1")
    
    # Fecha del reporte
    ws_resumen["A2"] = f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    if request.mes and request.año:
        meses = ["", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
                 "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
        ws_resumen["A3"] = f"Período: {meses[request.mes]} {request.año}"
    
    # Resumen de totales
    ws_resumen["A5"] = "RESUMEN GENERAL"
    ws_resumen["A5"].font = Font(bold=True, size=14)
    
    resumen_data = [
        ("Total Ganancias:", f"Bs. {total_ganancias:.2f}"),
        ("", ""),
        ("Ventas Lucio:", f"{len(ventas_lucio)} ventas"),
        ("- Total QR:", f"Bs. {total_qr:.2f}"),
        ("- Total Efectivo (Deuda):", f"Bs. {total_cash:.2f}"),
        ("", ""),
        ("Ventas Doctores:", f"{len(ventas_doctores)} ventas"),
        ("Ventas Envío:", f"{len(ventas_envio)} ventas"),
    ]
    
    for i, (label, value) in enumerate(resumen_data, start=7):
        ws_resumen[f"A{i}"] = label
        ws_resumen[f"B{i}"] = value
        if "Total Ganancias" in label:
            ws_resumen[f"A{i}"].font = Font(bold=True)
            ws_resumen[f"B{i}"].font = Font(bold=True)
    
    # Ajustar ancho de columnas
    ws_resumen.column_dimensions["A"].width = 25
    ws_resumen.column_dimensions["B"].width = 20
    
    # ==========================================
    # HOJA 2: Detalle Lucio
    # ==========================================
    ws_lucio = wb.create_sheet("Ventas Lucio")
    
    headers_lucio = ["Fecha", "Producto", "Cantidad", "Precio Venta", "Método Pago", "Cliente"]
    for col, header in enumerate(headers_lucio, start=1):
        cell = ws_lucio.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center")
    
    for row, venta in enumerate(ventas_lucio, start=2):
        ws_lucio.cell(row=row, column=1, value=venta.createdAt.strftime("%d/%m/%Y"))
        ws_lucio.cell(row=row, column=2, value=venta.productoNombre)
        ws_lucio.cell(row=row, column=3, value=venta.cantidad)
        ws_lucio.cell(row=row, column=4, value=f"Bs. {venta.precioVenta:.2f}")
        ws_lucio.cell(row=row, column=5, value=venta.metodoPago or "-")
        ws_lucio.cell(row=row, column=6, value=venta.clienteNombre or "-")
    
    # Ajustar anchos
    for col in range(1, 7):
        ws_lucio.column_dimensions[chr(64 + col)].width = 15
    
    # ==========================================
    # HOJA 3: Detalle Doctores
    # ==========================================
    ws_doctores = wb.create_sheet("Ventas Doctores")
    
    headers_doctores = ["Fecha", "Producto", "Cantidad", "Precio Venta", "Doctor"]
    for col, header in enumerate(headers_doctores, start=1):
        cell = ws_doctores.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center")
    
    for row, venta in enumerate(ventas_doctores, start=2):
        ws_doctores.cell(row=row, column=1, value=venta.createdAt.strftime("%d/%m/%Y"))
        ws_doctores.cell(row=row, column=2, value=venta.productoNombre)
        ws_doctores.cell(row=row, column=3, value=venta.cantidad)
        ws_doctores.cell(row=row, column=4, value=f"Bs. {venta.precioVenta:.2f}")
        ws_doctores.cell(row=row, column=5, value=venta.doctorNombre or "-")
    
    for col in range(1, 6):
        ws_doctores.column_dimensions[chr(64 + col)].width = 15
    
    # ==========================================
    # HOJA 4: Detalle Envíos
    # ==========================================
    ws_envio = wb.create_sheet("Ventas Envío")
    
    headers_envio = ["Fecha", "Producto", "Cantidad", "Precio Venta", "Departamento"]
    for col, header in enumerate(headers_envio, start=1):
        cell = ws_envio.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center")
    
    for row, venta in enumerate(ventas_envio, start=2):
        ws_envio.cell(row=row, column=1, value=venta.createdAt.strftime("%d/%m/%Y"))
        ws_envio.cell(row=row, column=2, value=venta.productoNombre)
        ws_envio.cell(row=row, column=3, value=venta.cantidad)
        ws_envio.cell(row=row, column=4, value=f"Bs. {venta.precioVenta:.2f}")
        ws_envio.cell(row=row, column=5, value=venta.departamento or "-")
    
    for col in range(1, 6):
        ws_envio.column_dimensions[chr(64 + col)].width = 15
    
    # Guardar a buffer de memoria
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    # Generar nombre del archivo
    if request.mes and request.año:
        filename = f"ganancias_{request.año}_{str(request.mes).zfill(2)}.xlsx"
    else:
        filename = f"ganancias_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    # Devolver como descarga
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
